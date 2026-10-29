"""UX1 full interactions through the real inbox/cards/store with synthetic ports."""
import copy
import io
import json
from datetime import datetime,timedelta
import pytest
from docx import Document
from fs_agent.scheduler import TZ
from fs_agent.project_catalog import digest
from test_am_self_fix1 import service,am_rig,base_rig,make_plan
from test_gateway import event,drain
from test_am_self_d import Sheet
from test_doc_logic import Business
from fs_agent.feishu import SendUnknown

@pytest.fixture
def ux(service):
    am,g,s,api,clock,port=service
    am.flags.update(attachments=True,precheck=True)
    am.updates.port=Sheet();business=Business()
    business.settings={'targets':{'A':{'label':'项目A / 测试节点 / 测试报告'},'B':{'label':'项目B / 交付节点 / 交付物'}}}
    g.logic.business=business
    return service

def latest(ux):return ux[2].rows("SELECT * FROM am_cards WHERE flow='ux' ORDER BY rowid DESC LIMIT 1")[0]

def request(ux,text,number,**metadata):
    am,g,s,api,clock,_=ux
    payload=event(text,number=number);payload['event']['message'].update(create_time=str(int(clock()*1000)),**metadata)
    assert am.ingest(payload)==(True,True);assert am.step();drain(g)
    return s.rows('SELECT * FROM am_inbox ORDER BY rowid DESC LIMIT 1')[0]

def click(ux,page,*,form=None,c=None,actor='ou_a',chat='oc_dm',overrides=None):
    am,g,s,api,clock,_=ux;c=c or latest(ux)
    actions=json.loads(c['snapshot'])['actions']
    match=next((key for key,value in actions.items() if value['page']==page and all(value.get(k)==v for k,v in (overrides or {}).items())),None)
    assert match is not None,(page,actions)
    n=len(s.rows('SELECT id FROM am_inbox'))
    payload={'header':{'app_id':g.config.app_id,'tenant_key':g.config.tenant_key,'event_id':f'ux_click_{n}','event_type':'card.action.trigger'},'event':{'operator':{'open_id':actor},'context':{'open_chat_id':chat,'open_message_id':c['message_id']},'action':{'tag':'button','value':{'flow':'am_self','nonce':c['nonce'],'action':match},'form_value':form}}}
    assert am.ingest(payload)==(True,True);assert am.step();drain(g)
    return s.rows('SELECT * FROM am_inbox ORDER BY rowid DESC LIMIT 1')[0]

def text(ux):return ux[3].sent[-1]['body']

def reminder_form(ux):
    request(ux,'/menu',1);click(ux,'projects',overrides={'next':'remind'});click(ux,'remind');click(ux,'remind')
    return latest(ux)

def values(ux,**kw):
    when=(datetime.fromtimestamp(ux[4](),TZ)+timedelta(minutes=5)).strftime('%Y-%m-%d %H:%M +0800')
    return dict({'people':['0'],'body':'请提交完整报告，尚未验收。','when':when,'frequency':'once','until':'','policy':'preview_each'},**kw)

def test_UX01_menu_event_to_live_project_detail(ux):
    am,g,s,api,clock,_=ux
    menu={'header':{'app_id':g.config.app_id,'tenant_key':g.config.tenant_key,'event_id':'menu_1','event_type':'application.bot.menu_v6'},'event':{'operator':{'operator_id':{'open_id':'ou_a'}},'event_key':'am_work','timestamp':int(clock()*1000)}}
    assert am.ingest(menu)==(True,True);assert am.ingest(menu)==(True,False)
    am.step();drain(g)
    c=latest(ux)
    g.reminders.port.read=lambda mid:{'chat_id':'oc_dm','sender_id':g.config.bot_open_id}
    click(ux,'projects',c=c,overrides={})
    click(ux,'project')
    assert '需求75' in text(ux) and '实际收货尚未核实' in text(ux)
    assert len(s.rows('SELECT * FROM action_items'))==0

@pytest.mark.parametrize('actor,key',[('ou_other','am_work'),('ou_a','/plan 任意指令')])
def test_menu_rejects_untrusted_actor_or_event_key(ux,actor,key):
    am,g,s,api,clock,_=ux
    p={'header':{'app_id':g.config.app_id,'tenant_key':g.config.tenant_key,'event_id':'bad','event_type':'application.bot.menu_v6'},'event':{'operator':{'operator_id':{'open_id':actor}},'event_key':key,'timestamp':int(clock()*1000)}}
    assert am.ingest(p)==(True,False) and not s.rows('SELECT * FROM am_inbox')

def test_UX02_edit_preserves_form_and_invalidates_old_approval(ux):
    form=reminder_form(ux);click(ux,'preview_reminder',form=values(ux))
    preview=latest(ux);assert '完整提醒预览' in text(ux) and 'Asia/Shanghai' in text(ux)
    assert not ux[2].rows('SELECT * FROM reminder_plans')
    click(ux,'remind');edited=latest(ux)
    assert json.loads(edited['snapshot'])['fields']['body']['default']==values(ux)['body']
    r=click(ux,'confirm_reminder',c=preview)
    assert r['error']=='UX_DRAFT_CHANGED' and not ux[2].rows('SELECT * FROM reminder_plans')
    click(ux,'preview_reminder',c=edited,form=values(ux,body='修改后的完整提醒'))
    click(ux,'confirm_reminder');assert len(ux[2].rows('SELECT * FROM reminder_plans'))==1

def test_UX03_confirmation_replay_creates_one_plan_and_item(ux):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux));preview=latest(ux)
    click(ux,'confirm_reminder');click(ux,'confirm_reminder',c=preview)
    assert len(ux[2].rows('SELECT * FROM reminder_plans'))==1
    assert len(ux[2].rows('SELECT * FROM action_items'))==1
    assert not ux[2].rows('SELECT * FROM reminder_occurrences')

def test_UX03_unknown_only_offers_reconciliation(ux):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux));click(ux,'confirm_reminder')
    am,g,s,api,clock,_=ux;p=s.rows('SELECT * FROM reminder_plans')[0]
    clock.now=p['next_at'];am.scheduler.scan();o=s.rows('SELECT * FROM reminder_occurrences')[0]
    s.execute("UPDATE reminder_occurrences SET state='unknown',readback_state='pending' WHERE id=?",(o['id'],))
    click(ux,'plan');snapshot=json.loads(latest(ux)['snapshot'])
    assert '结果待核对' in text(ux)
    assert any(a['page']=='reconcile' for a in snapshot['actions'].values())
    assert not any('retry' in a['page'] or 'send' in a['page'] for a in snapshot['actions'].values())

@pytest.mark.parametrize('change',[{'people':['999']},{'body':''},{'when':'明天'},{'frequency':'workdays','until':'2033-06-01 09:00 +0800'}])
def test_form_invalid_values_never_create_plan(ux,change):
    reminder_form(ux);r=click(ux,'preview_reminder',form=values(ux,**change))
    assert r['state']=='failed' and not ux[2].rows('SELECT * FROM reminder_plans')

def test_UX05_original_progress_message_fingerprint(ux):
    request(ux,'/menu',1);click(ux,'projects',overrides={'next':'collect'});click(ux,'collect')
    message='  报告明天提交，测试尚未完成。  '
    r=request(ux,message,2)
    record=ux[2].rows('SELECT * FROM progress_collections')[0]
    source=json.loads(record['source'])
    assert source['message_id']=='om2' and source['content_fingerprint']==digest({'text':message})
    assert len(ux[2].rows('SELECT * FROM ledger_proposals'))==1 and not ux[1].logic.business.writes
    assert not ux[0].updates.port.calls and ux[0].flags['ledger_write'] is False

def test_collection_expiry_and_switch_never_rebind_queued_message(ux):
    request(ux,'/menu',1);click(ux,'projects',overrides={'next':'collect'});click(ux,'collect')
    am,g,s,api,clock,_=ux
    payload=event('样机完成，待核验',number=2);payload['event']['message']['create_time']=str(int(clock()*1000))
    assert am.ingest(payload)==(True,True)
    ctx=s.rows('SELECT * FROM am_ux_collection')[0];s.execute("UPDATE am_ux_collection SET token='changed' WHERE scope=?",(ctx['scope'],))
    am.step();drain(g)
    assert not s.rows('SELECT * FROM progress_collections')

def test_UX06_precheck_file_bound_to_exact_target_and_F08(ux):
    request(ux,'/menu',1);click(ux,'targets');click(ux,'precheck_form')
    click(ux,'start_precheck',form={'kind':'test','requirements':'交付物清单'})
    am,g,s,api,clock,_=ux
    doc=Document();t=doc.add_table(rows=1,cols=4)
    for i,word in enumerate(('交付物清单','','负责人','AR001合成人员')):t.cell(0,i).text=word
    output=io.BytesIO();doc.save(output);api.read_file=lambda *a,**k:output.getvalue()
    p=event('',number=2);p['event']['message'].update(message_type='file',content=json.dumps({'file_key':'synthetic','file_name':'合成.docx'}),create_time=str(int(clock()*1000)))
    assert am.ingest(p)==(True,True);am.step();drain(g)
    d=s.rows('SELECT * FROM logic_drafts ORDER BY rowid DESC')[0]
    assert d['state']=='ready',s.rows('SELECT state,error FROM am_inbox')
    assert json.loads(d['target'])['id']=='A'
    report=json.loads(d['analysis'])['categories']
    assert not report['已具备'] and report['缺少'] and not g.logic.business.writes

@pytest.mark.parametrize('actor,chat',[('ou_b','oc_dm'),('ou_a','oc_other')])
def test_UX07_wrong_actor_or_chat_cannot_execute(ux,actor,chat):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux));preview=latest(ux)
    result=click(ux,'confirm_reminder',c=preview,actor=actor,chat=chat)
    assert result['state']=='failed' and not ux[2].rows('SELECT * FROM reminder_plans')

def test_UX07_expired_card_cannot_execute(ux):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux));preview=latest(ux)
    ux[4].now+=3601
    result=click(ux,'confirm_reminder',c=preview)
    assert result['error']=='CARD_IDENTITY_OR_VERSION_INVALID'
    assert not ux[2].rows('SELECT * FROM reminder_plans')

def test_UX04_group_form_continuous_feedback_two_actors(ux):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux,people=['0','1']));click(ux,'confirm_reminder')
    am,g,s,api,clock,_=ux;p=s.rows('SELECT * FROM reminder_plans')[0]
    clock.now=p['next_at'];am.scheduler.scan();o=s.rows('SELECT * FROM reminder_occurrences')[0]
    saved=json.loads(o['payload']);s.execute("UPDATE reminder_occurrences SET state='sent',message_id='om_feedback',readback_state='verified' WHERE id=?",(o['id'],))
    def feedback(actor,action,form=None):
        n=len(s.rows('SELECT * FROM am_inbox'))
        payload={'header':{'app_id':g.config.app_id,'tenant_key':g.config.tenant_key,'event_id':f'f{n}','event_type':'card.action.trigger'},'event':{'operator':{'open_id':actor},'context':{'open_chat_id':'oc_reminder','open_message_id':'om_feedback'},'action':{'value':{'flow':'am_self','feedback':o['id'],'nonce':saved['feedback_nonce'],'action':action},'form_value':form}}}
        assert am.ingest(payload)==(True,True);am.step();drain(g)
    feedback('ou_engineer_a','received')
    feedback('ou_engineer_a','feedback_form',{'kind':'eta','eta':values(ux)['when'],'detail':'预计明天完成'})
    feedback('ou_engineer_a','feedback_form',{'kind':'submitted','detail':'https://synthetic.invalid/A'})
    feedback('ou_engineer_b','feedback_form',{'kind':'submitted','detail':'https://synthetic.invalid/B'})
    history=am.items.history(o['item_id']);assert len(history)==4 and all(h['applied'] for h in history)
    assert 'synthetic.invalid/A' in am.item_text(am.items.get(o['item_id']))
    assert any(x['user']=='ou_engineer_b' and '待核验' in x['body'] for x in api.sent)

def test_UX08_search_no_results_and_long_body_readable(ux):
    request(ux,'/menu',1);click(ux,'projects');click(ux,'search',form={'query':'不会存在的项目'})
    assert '未找到匹配项目' in text(ux)
    click(ux,'home');click(ux,'projects',overrides={'next':'remind'});click(ux,'remind');click(ux,'remind')
    body='长正文。'*240;click(ux,'preview_reminder',form=values(ux,body=body))
    sent=json.loads(text(ux));alltext=''.join(e.get('text',{}).get('content','') for e in sent['body']['elements'])
    assert body in alltext and all(len(e.get('text',{}).get('content',''))<=3500 for e in sent['body']['elements'] if e['tag']=='div')

def test_collection_cross_thread_requires_new_choice(ux):
    request(ux,'/menu',1,thread_id='thread-a');click(ux,'projects',overrides={'next':'collect'});click(ux,'collect')
    result=request(ux,'这里是另一话题的材料',2,thread_id='thread-b')
    assert result['error']=='COLLECTION_CONTEXT_EXPIRED'
    assert not ux[2].rows('SELECT * FROM progress_collections')

def test_collection_different_project_mention_rejected(ux):
    request(ux,'/menu',1);click(ux,'projects',overrides={'next':'collect'});click(ux,'collect')
    result=request(ux,'需求76的测试完成了',2)
    assert result['error']=='COLLECTION_REFERENCE_CONFLICT' and not ux[2].rows('SELECT * FROM progress_collections')

def test_subscription_edit_and_repeat_confirm(ux):
    request(ux,'/menu',1);click(ux,'work');click(ux,'more');click(ux,'subscribe')
    until=(datetime.fromtimestamp(ux[4](),TZ)+timedelta(days=3)).strftime('%Y-%m-%d %H:%M +0800')
    v={'frequency':'daily','time':'09:00','until':until}
    click(ux,'subscription_preview',form=v);preview=latest(ux)
    click(ux,'subscribe');edited=latest(ux)
    assert json.loads(edited['snapshot'])['fields']['until']['default']==until
    assert click(ux,'subscription_confirm',c=preview)['error']=='UX_DRAFT_CHANGED'
    click(ux,'subscription_preview',form=v,c=edited);preview=latest(ux)
    click(ux,'subscription_confirm');click(ux,'subscription_confirm',c=preview)
    assert len(ux[2].rows('SELECT * FROM am_subscriptions'))==1
    assert len(ux[2].rows('SELECT * FROM reminder_plans'))==1

def test_future_cycle_start_respected(ux):
    reminder_form(ux)
    when=datetime.fromtimestamp(ux[4](),TZ)+timedelta(days=3)
    v=values(ux,frequency='exact_daily',when=when.strftime('%Y-%m-%d %H:%M +0800'),until=(when+timedelta(days=3)).strftime('%Y-%m-%d %H:%M +0800'))
    click(ux,'preview_reminder',form=v);click(ux,'confirm_reminder')
    p=ux[2].rows('SELECT * FROM reminder_plans')[0]
    assert p['next_at']>=when.timestamp()-60

def test_old_form_is_blocked_after_feature_disabled(ux):
    reminder_form(ux);click(ux,'preview_reminder',form=values(ux));preview=latest(ux)
    ux[0].flags['scheduler']=False
    assert click(ux,'confirm_reminder',c=preview)['error']=='FEATURE_NOT_ENABLED'
    assert not ux[2].rows('SELECT * FROM reminder_plans')

def test_precheck_unread_file_is_not_claimed_missing(ux):
    from fs_agent.config import BoundaryError
    request(ux,'/menu',1);click(ux,'targets');click(ux,'precheck_form');click(ux,'start_precheck',form={'kind':'test','requirements':'交付物清单'})
    def unread(*a,**kw):raise BoundaryError('FILE_READ_FAILED')
    am,g,s,api,clock,_=ux;api.read_file=unread
    p=event('',number=2);p['event']['message'].update(message_type='file',content=json.dumps({'file_key':'synthetic','file_name':'不可读.docx'}),create_time=str(int(clock()*1000)))
    assert am.ingest(p)==(True,True);am.step();drain(g)
    d=s.rows('SELECT * FROM logic_drafts')[0];report=json.loads(d['analysis'])['categories']
    assert not report['缺少'] and report['需专业判断']

def test_project_card_reply_binding_survives_later_query(ux):
    request(ux,'/menu',1);click(ux,'projects');click(ux,'project')
    reply=latest(ux)['message_id']
    request(ux,'/project 需求76',2)
    request(ux,'那谁负责？',3,parent_id=reply)
    assert 'AR1测试甲' in text(ux) and 'AR2乙' not in text(ux)

def test_datetime_form_uses_explicit_business_timezone_not_client_reference():
    from am_ux import card_time
    from fs_agent.card_ui import field
    assert card_time('2026-10-30 09:00 +0000')==card_time('2026-10-30 09:00 +0800')
    assert field('when','北京时间',kind='picker_datetime',default='2026-10-30 09:00 +0000')[-1]['initial_datetime']=='2026-10-30 09:00'

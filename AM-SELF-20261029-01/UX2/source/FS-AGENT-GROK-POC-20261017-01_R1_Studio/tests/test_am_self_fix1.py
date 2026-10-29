"""F01–F08 assert correct behavior through real SQLite and synthetic ports only."""
import copy
import io
import json
import sys
from contextlib import closing
from pathlib import Path
from datetime import datetime
from types import SimpleNamespace
import threading
from concurrent.futures import ThreadPoolExecutor
import pytest

sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'local-qwen3.8/feishu-bridge'))
from fs_agent.action_items import ActionItems
from fs_agent.scheduler import Scheduler,TZ
from fs_agent.store import Store
from fs_agent.feishu import SendUnknown,SendRejected
from fs_agent.doc_logic import digest,read_source
from fs_agent.precheck import inspect_materials
from fs_agent.config import BoundaryError
from test_gateway import rig as base_rig,event,drain
from test_reminder import rig as am_rig
from test_doc_logic import rig as logic_rig,ready
from test_meegle_port import port_fixture
from test_am_self_c import Port as SendPort
from test_group_ledger import Port as ReadPort,VALUES
from group_ledger import GroupLedger
from am_service import AmSelfService

@pytest.fixture
def service(am_rig,tmp_path):
    g,s,api,model,clock,e=am_rig
    port=ReadPort();port.authorize_user=lambda user:None
    port.values=copy.deepcopy(VALUES[:2])+[['需求76:乙项目','','','跟进乙报告','乙项目处理中','AR2乙','']]
    (tmp_path/'ledger').mkdir()
    ledger=GroupLedger(tmp_path/'ledger',port,lambda *a:pytest.fail('no group sends'),clock=clock)
    am=AmSelfService(g,ledger,flags={'query':True,'snapshots':True,'items':True,'scheduler':True,'progress_collection':True})
    yield am,g,s,api,clock,port
    ledger.close()

def ask(rig,text,number,**metadata):
    am,g,s,api,clock,port=rig
    payload=event(text,number=number);payload['event']['message'].update(metadata,create_time=str(int(clock()*1000)))
    assert am.ingest(payload)==(True,True)
    assert am.step();drain(g)
    return api.sent[-1]['body']

def make_plan(sch,now,item,op,kind='once'):
    rule={'at':now+60} if kind=='once' else {'time':'09:00','timezone':'Asia/Shanghai'}
    p=sch.create(sch.owner,op,item_id=item,kind=kind,rule=rule,
                 target={'kind':'group','payload':{'synthetic':True}},template='请提交报告。',policy='fixed_rule',valid_until=now+3*86400)
    return sch.approve_plan(sch.owner,p['id'],p['version'],p['hash'])

def test_F01_queued_duplicate_unknown_does_not_block_independent_item(tmp_path):
    now=[datetime(2026,10,29,8,tzinfo=TZ).timestamp()]
    with closing(Store(tmp_path/'state.db',lambda:now[0])) as s:
        items=ActionItems(s,'ou_owner');port=SendPort();sch=Scheduler(s,items,'ou_owner',port,enabled=lambda:True)
        def item(op):return items.create('ou_owner',op,project_ref='project',description='报告',requirements='核验',assignees=['ou_a'],source={})['id']
        a,b=item('a'),item('b')
        for op,ident in [('a1',a),('a2',a),('b1',b)]:make_plan(sch,now[0],ident,op)
        now[0]+=60;sch.scan();queued=s.rows('SELECT * FROM reminder_occurrences ORDER BY rowid')
        port.error=SendUnknown('lost');sch.dispatch();port.error=None
        sch.dispatch()
        assert port.calls==[queued[0]['request_key'],queued[2]['request_key']]
        assert sch.occurrence(queued[1]['id'])['state']=='ready'
        sch=Scheduler(s,items,'ou_owner',port,enabled=lambda:True)
        assert not sch.dispatch()
        now[0]+=61;sch.dispatch()
        assert sch.occurrence(queued[1]['id'])['state']=='cancelled'
        assert len(port.calls)==2

def test_F02_same_card_continuous_multi_person_feedback_and_owner_fence(service):
    am,g,s,api,clock,port=service
    item=am.items.create(am.owner,'item',project_ref='p',description='交报告',requirements='核验材料',assignees=['ou_engineer_a','ou_engineer_b'],source={})
    results=[]
    for actor,op,kind,body in [('ou_engineer_a','ack','received','收到'),('ou_engineer_a','eta','eta','明天'),('ou_engineer_a','submit-a','submitted','https://synthetic.invalid/report-A'),('ou_engineer_b','submit-b','submitted','https://synthetic.invalid/report-B')]:
        result=am.items.feedback(actor,op,item['id'],1,kind,body,{'message_id':'om_card','submission_message':op})
        results.append(result)
    assert all(r['applied'] for r in results)
    assert results[0]['item']['state']=='waiting' and results[-1]['item']['state']=='submitted'
    text=am.item_text(results[-1]['item'])
    assert 'report-A' in text and 'report-B' in text and 'om_card' in text
    assert am.items.feedback('ou_engineer_b','submit-b',item['id'],1,'submitted','https://synthetic.invalid/report-B',{'message_id':'om_card','submission_message':'submit-b'})==results[-1]
    changed=am.items.change(am.owner,'pause',item['id'],results[-1]['item']['version'],'pause')
    late=am.items.feedback('ou_engineer_a','late',item['id'],1,'help','迟到协助',{})
    assert not late['applied'] and late['item']['version']==changed['version']
    assert '迟到协助' in am.item_text(late['item'])

def test_F03_private_reference_overrides_recent_focus(service):
    ask(service,'/project 需求75',1)
    first_reply='om_'+service[2].rows('SELECT id FROM outbox ORDER BY rowid DESC LIMIT 1')[0]['id']
    ask(service,'/project 需求76',2)
    result=ask(service,'那谁负责？',3,parent_id='om1')
    assert 'AR1测试甲' in result and 'AR2乙' not in result
    result=ask(service,'那谁负责？',4,parent_id=first_reply)
    assert 'AR1测试甲' in result and 'AR2乙' not in result
    result=ask(service,'那谁负责？',5,parent_id='om_missing')
    assert '需求76' not in result and '确认' in result

def briefing_plan(service):
    am,g,s,api,clock,port=service
    clock.now=datetime(2026,10,29,8,tzinfo=TZ).timestamp()
    row={'user':am.owner};am.projects(row)
    port.values[1][4]='自报新变化：报告A待核验';am.projects(row)
    p=am.scheduler.create(am.owner,'briefing',item_id=None,kind='briefing',rule={'time':'09:00','timezone':'Asia/Shanghai'},target={'kind':'owner','user':am.owner,'chat':'oc_dm'},template='',policy='owner_subscription',valid_until=clock()+3*86400)
    am.scheduler.approve_plan(am.owner,p['id'],1,p['hash'])
    s.execute("INSERT INTO am_subscriptions VALUES('sub',?,?,'active')",(am.owner,p['id']))
    transport=SendPort();am.scheduler.port=transport;am.scheduler.notify=lambda *a:None
    clock.now+=3600;am.scheduler.scan()
    return p,transport

@pytest.mark.parametrize('outcome',['preflight_expiry','rejected_exhausted','unknown','verified'])
def test_F04_next_briefing_body_and_attempt_association(service,outcome):
    am,g,s,api,clock,port=service;p,transport=briefing_plan(service)
    first=s.rows('SELECT * FROM reminder_occurrences')[0]
    if outcome=='preflight_expiry':transport.before=lambda:setattr(clock,'now',clock()+61)
    elif outcome=='rejected_exhausted':transport.error=SendRejected('definitely not sent')
    elif outcome=='unknown':transport.error=SendUnknown('lost')
    else:transport.found=True
    am.scheduler.dispatch()
    if outcome=='rejected_exhausted':
        for _ in range(3):clock.now+=10;am.scheduler.dispatch()
    transport.before=lambda:None;transport.error=None
    clock.now=datetime(2026,10,30,9,tzinfo=TZ).timestamp();am.scheduler.scan()
    occurrences=s.rows('SELECT * FROM reminder_occurrences ORDER BY rowid')
    if outcome=='unknown':
        assert len(occurrences)==1 and len(transport.calls)==1
    elif outcome=='verified':
        assert '报告A待核验' not in json.loads(occurrences[-1]['payload'])['body']
    else:
        assert len(occurrences)==2
        assert '报告A待核验' in json.dumps(json.loads(occurrences[-1]['payload'])['briefing_data']['events'],ensure_ascii=False)
        delivery=s.rows('SELECT * FROM am_briefing_deliveries')[0]
        assert delivery['occurrence_id']==occurrences[-1]['id']
        attempts=s.rows('SELECT * FROM am_briefing_delivery_attempts ORDER BY rowid')
        assert [a['state'] for a in attempts]==['not_sent','pending']
        assert am.scheduler.occurrence(first['id'])['state']=='cancelled'
        transport.found=True;am.scheduler.dispatch()
        assert s.rows('SELECT state FROM am_briefing_deliveries')[0]['state']=='verified'

def test_F05_briefing_uses_same_current_set_as_todo(service):
    ask(service,'/todo',1)
    service[-1].values=service[-1].values[:2]
    assert '跟进乙报告' not in ask(service,'/todo',2)
    assert '跟进乙报告' not in ask(service,'/briefing',3)
    assert service[0].ledger.db.execute('SELECT count(*) FROM project_snapshots').fetchone()[0]==2

def prechecked(logic_rig):
    g,s,row,d=ready(logic_rig);port,state,settings=port_fixture(g);g.logic.business=port
    state['files']=[{'name':'original.md','type':'text/markdown','size':'8','fileToken':'original-token','url':'https://files.invalid/original.md'}]
    state['data']['original.md']=b'original'
    d=g.logic.start_precheck(row,'report','test');g.logic.set_precheck_requirements(row,'交付物清单')
    source=read_source('report.txt','交付物清单：\n1. 测试报告A'.encode(),source_id='S1',origin='synthetic')
    sources=[source]+port.sources(row,json.loads(d['target']))
    s.execute('UPDATE logic_drafts SET sources=? WHERE id=?',(json.dumps(sources),d['id']))
    d=g.logic.get(row,d['id']);g.logic.run_precheck(dict(row,id='synthetic-task',run_id='run'),d,sources,{'revision':d['revision']})
    return g,s,row,g.logic.get(row,d['id']),port,state,settings

def test_F06_precheck_adopt_actual_adapter_mount_and_readback(logic_rig):
    g,s,row,d,port,state,settings=prechecked(logic_rig)
    assert d['content_hash']==digest(d['content'].encode())
    g.logic.adopt(row,d['id']);result=g.logic.mount(row,d['id'])
    assert result['receipt']['previous_tokens']==['original-token']
    assert state['uploads']==state['updates']==1 and len(state['files'])==2
    assert state['review']=='not_reviewed'
    assert g.logic.mount(row,d['id'])['receipt']==result['receipt']
    assert state['uploads']==1

@pytest.mark.parametrize('kind',['text','post','whitespace'])
@pytest.mark.parametrize('correction',[False,True])
def test_F07_source_adapter_raw_content_preserved(service,kind,correction):
    am,g,s,api,clock,port=service
    q='/update 需求75：报告预计明天提交'
    if correction:
        ask(service,q,99)
        r=s.rows('SELECT * FROM progress_collections')[0]
        q=f"/update 更正 {r['id']} v{r['version']}：报告预计后天提交"
    content={'text':q} if kind=='text' else {'text':' \n'+q+'\n '} if kind=='whitespace' else {'title':'','content':[[{'tag':'text','text':q}]]}
    ask(service,q,1,message_type='post' if kind=='post' else 'text',content=json.dumps(content))
    collection=s.rows('SELECT * FROM progress_collections ORDER BY rowid DESC LIMIT 1')[0]
    source=json.loads(collection['source'])
    remote={'content':json.dumps(content),'deleted':False,'sender_id':am.owner,'sender_type':'user','chat_id':'oc_dm','msg_type':'post' if kind=='post' else 'text'}
    g.reminders.port.read=lambda mid:remote
    assert am.updates.port.source_current(source)
    remote['content']=json.dumps({'text':'changed'})
    assert not am.updates.port.source_current(source)
    remote['content']=json.dumps(content);remote['deleted']=True
    assert not am.updates.port.source_current(source)
    assert am.flags['ledger_write'] is False

@pytest.mark.parametrize('shape',['text','docx','table','positive','partial'])
def test_F08_empty_heading_and_table_are_not_substantive(shape):
    text='交付物清单：\n   '
    if shape=='positive':text='交付物清单：\n1. 测试报告A'
    if shape in ('docx','table'):
        from docx import Document
        doc=Document()
        if shape=='docx':doc.add_heading('交付物清单',1)
        else:
            table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='交付物清单'
        buf=io.BytesIO();doc.save(buf);source=read_source('report.docx',buf.getvalue(),source_id='S1',origin='synthetic')
    else:source=read_source('report.txt',text.encode(),source_id='S1',origin='synthetic')
    if shape=='partial':source.update(read_status='partial',limitation='部分内容不可读')
    report=inspect_materials([source],{'checks':[{'id':'1','label':'交付物清单','text':'交付物清单','kind':'presence'}]})
    if shape=='positive':assert report['已具备'] and not report['缺少']
    elif shape=='partial':assert not report['已具备'] and not report['缺少'] and report['需专业判断']
    else:assert not report['已具备'] and report['缺少']


@pytest.mark.parametrize('phase',['claim','guard'])
def test_F01_blocker_added_during_preflight_or_before_send(service,phase):
    am,g,s,api,clock,read=service
    item=am.items.create(am.owner,'race-item',project_ref='p',description='报告',requirements='核验',assignees=['ou_a'],source={})
    for op in ('first','second'):make_plan(am.scheduler,clock(),item['id'],op)
    transport=SendPort();am.scheduler.port=transport;am.scheduler.notify=lambda *a:None
    clock.now+=60;am.scheduler.scan();a,b=s.rows('SELECT * FROM reminder_occurrences ORDER BY rowid')
    def block():s.execute("UPDATE reminder_occurrences SET state='sent',readback_state=NULL WHERE id=?",(b['id'],))
    if phase=='claim':transport.before=block
    else:
        original=transport.send
        def send(payload,key,guard):block();return original(payload,key,guard)
        transport.send=send
    am.scheduler.dispatch()
    assert not transport.calls
    assert am.scheduler.occurrence(a['id'])['request_key']==a['request_key']
    assert am.scheduler.occurrence(a['id'])['deadline']==a['deadline']


def test_F01_concurrent_claims_same_followup_have_one_external_call(service):
    am,g,s,api,clock,read=service
    item=am.items.create(am.owner,'race-item',project_ref='p',description='报告',requirements='核验',assignees=['ou_a'],source={})
    for op in ('first','second'):make_plan(am.scheduler,clock(),item['id'],op)
    transport=SendPort();am.scheduler.port=transport;am.scheduler.notify=lambda *a:None
    clock.now+=60;am.scheduler.scan();occurrences=s.rows('SELECT * FROM reminder_occurrences ORDER BY rowid')
    barrier=threading.Barrier(2);transport.before=lambda:barrier.wait(3);transport.error=SendUnknown('lost')
    with ThreadPoolExecutor(2) as pool:list(pool.map(am.scheduler.dispatch_one,occurrences))
    assert len(transport.calls)==1


def test_F02_actual_feedback_binding_and_closed_old_card(service):
    am,g,s,api,clock,read=service
    item=am.items.create(am.owner,'bound-item',project_ref='p',description='报告',requirements='核验',assignees=['ou_engineer_a','ou_engineer_b'],source={})
    p=make_plan(am.scheduler,clock(),item['id'],'bound-plan')
    # An actual BusinessPort card is prepared from the existing synthetic target.
    from fs_agent.reminder_content import normalized_payload
    target=normalized_payload(item['id'],g.reminders.port.bot_info,g.reminders.port.group_info,g.reminders.port.people,'报告')
    s.execute('UPDATE reminder_plans SET target=? WHERE id=?',(json.dumps({'kind':'group','payload':target}),p['id']))
    clock.now+=60;am.scheduler.scan();o=s.rows('SELECT * FROM reminder_occurrences')[0]
    s.execute("UPDATE reminder_occurrences SET state='sent',message_id='om_card',readback_state='verified' WHERE id=?",(o['id'],))
    nonce=json.loads(o['payload'])['feedback_nonce']
    def feedback(actor,kind,detail,ident):
        row={'id':ident,'user':actor,'chat':'oc_reminder','app':g.config.app_id,'tenant':g.config.tenant_key,'chat_type':'p2p','expires':clock()+86400}
        body={'value':{'feedback':o['id'],'nonce':nonce,'action':kind},'source_message':'om_card','detail':detail}
        am.handle_feedback(row,body)
    feedback('ou_engineer_a','received','收到','ack')
    feedback('ou_engineer_a','submitted','材料A','a')
    feedback('ou_engineer_b','submitted','材料B','b')
    feedback('ou_engineer_b','submitted','材料B','retry')
    assert len(am.items.history(item['id']))==3 and all(h['applied'] for h in am.items.history(item['id']))
    current=am.items.get(item['id']);am.items.change(am.owner,'close-bound',item['id'],current['version'],'close',evidence='合成材料人工核验')
    feedback('ou_engineer_a','submitted','迟到材料','late')
    assert am.items.get(item['id'])['state']=='closed' and not am.items.history(item['id'])[-1]['applied']


def test_F03_thread_identity_expiry_and_ttl_cleanup(service):
    am,g,s,api,clock,read=service
    ask(service,'/project 需求75',1,thread_id='thread-a')
    ask(service,'/project 需求76',2,thread_id='thread-b')
    assert 'AR1测试甲' in ask(service,'那谁负责？',3,thread_id='thread-a')
    assert 'AR2乙' in ask(service,'那谁负责？',4,thread_id='thread-b')
    assert '确认' in ask(service,'那谁负责？',5,thread_id='thread-b',parent_id='om1')
    ask(service,'/project 需求76',6)
    s.execute("INSERT INTO am_query_refs VALUES('om_other',?,? ,?)",(json.dumps(['tenant_test','cli_test','oc_dm','ou_other'],separators=(',',':')),s.rows("SELECT project_ref FROM am_query_refs WHERE message_id='om1'")[0]['project_ref'],clock()+100))
    assert '确认' in ask(service,'那谁负责？',7,parent_id='om_other')
    clock.now+=86401
    assert '确认' in ask(service,'那谁负责？',8,parent_id='om6')
    assert not s.rows("SELECT * FROM am_source_meta WHERE id='om6'")
    assert not s.rows("SELECT * FROM am_query_refs WHERE message_id='om6'")


def test_F06_local_failure_is_definite_but_upload_failure_stays_unknown(logic_rig):
    g,s,row,d,port,state,settings=prechecked(logic_rig)
    g.logic.adopt(row,d['id'])
    s.execute('UPDATE logic_drafts SET content=? WHERE id=?',(d['content']+'\n未采用的变更',d['id']))
    result=g.logic.mount(row,d['id'])
    assert '确定失败' in result['message'] and 'BUSINESS_CONTENT_CHANGED' in result['message'] and state['uploads']==0
    assert s.rows('SELECT state FROM logic_writes')[-1]['state']=='confirmed_not_sent'
    assert '确定失败' in g.logic.mount(row,d['id'])['message']
    assert '确定失败' in g.logic.reconcile_draft(row,d['id'])
    s.execute('UPDATE logic_drafts SET content=?,revision=revision+1 WHERE id=?',(d['content'],d['id']))
    run=port._run
    def lose(args,**kwargs):
        if args[:2]==['attachment','+upload']:
            run(args,**kwargs);raise TimeoutError('response lost')
        return run(args,**kwargs)
    port._run=lose
    assert '待确认' in g.logic.mount(row,d['id'])['message']
    assert state['uploads']==1 and state['updates']==0
    assert s.rows('SELECT state FROM logic_writes ORDER BY rowid')[-1]['state']=='unknown'


def test_F06_legacy_hash_migration_invalidates_adoption_preserves_external_unknown(logic_rig):
    from fs_agent.project_catalog import digest as structured_digest
    from fs_agent.logic_service import LogicService
    g,s,row,d,port,state,settings=prechecked(logic_rig)
    wrong=structured_digest(d['content'])
    s.execute("UPDATE logic_drafts SET content_hash=?,accepted_hash=? WHERE id=?",(wrong,wrong,d['id']))
    for ident,receipt in [('local',None),('external',json.dumps({'phase':'uploaded','upload':{'file_token':'synthetic'}}))]:
        s.execute("INSERT INTO logic_writes VALUES(?,?,?,?,?,'unknown',?)",(ident,d['id'],d['revision'],d['target'],wrong,receipt))
    s.execute("DELETE FROM meta WHERE key='R1-FIX1-precheck-hash-1'")
    g.logic=LogicService(g,business=port)
    changed=g.logic.get(row,d['id'])
    assert changed['state']=='collecting' and changed['accepted_hash'] is None and changed['content_hash']==wrong
    assert s.rows("SELECT state FROM logic_writes WHERE idem='local'")[0]['state']=='confirmed_not_sent'
    assert s.rows("SELECT state FROM logic_writes WHERE idem='external'")[0]['state']=='unknown'
    g.logic=LogicService(g,business=port)
    assert g.logic.get(row,d['id'])['revision']==changed['revision']
    assert state['uploads']==state['updates']==0


def test_F07_legacy_unproven_fingerprint_cannot_authorize_source(service):
    from fs_agent.project_catalog import digest as structured_digest
    am,g,s,api,clock,read=service
    remote={'deleted':False,'content':{'text':'normalized'}}
    g.reminders.port.read=lambda _:remote
    assert not am.updates.port.source_current({'message_id':'om_legacy','content_fingerprint':structured_digest(remote['content'])})


def test_F02_legacy_item_migration_preserves_decision_fence_and_history(service):
    am,g,s,api,clock,read=service
    item=am.items.create(am.owner,'migrate-item',project_ref='p',description='报告',requirements='原要求',assignees=['ou_a'],source={})
    am.items.feedback('ou_a','migrate-ack',item['id'],1,'received','收到',{})
    s.execute('ALTER TABLE action_items DROP COLUMN decision_version')
    migrated=ActionItems(s,am.owner)
    assert migrated.get(item['id'])['decision_version']==1
    assert migrated.feedback('ou_a','migrate-submit',item['id'],1,'submitted','材料A',{})['applied']
    current=migrated.get(item['id']);migrated.change(am.owner,'migrate-check',item['id'],current['version'],'check',next_check=clock()+60)
    decided=migrated.get(item['id'])
    s.execute('ALTER TABLE action_items DROP COLUMN decision_version')
    migrated=ActionItems(s,am.owner)
    assert migrated.get(item['id'])['decision_version']==decided['version']
    assert not migrated.feedback('ou_a','migrate-late',item['id'],1,'submitted','迟到材料',{})['applied']
    counts={t:s.rows('SELECT count(*) AS n FROM '+t)[0]['n'] for t in ('action_items','item_feedback','item_operations')}
    ActionItems(s,am.owner)
    assert counts=={t:s.rows('SELECT count(*) AS n FROM '+t)[0]['n'] for t in counts}


@pytest.mark.parametrize('outcome',['cancelled','unknown','sent'])
def test_F04_legacy_pending_migration_retains_attempt_identity(service,outcome):
    am,g,s,api,clock,read=service;p,transport=briefing_plan(service)
    o=s.rows('SELECT * FROM reminder_occurrences')[0]
    s.execute('UPDATE reminder_occurrences SET state=?,readback_state=? WHERE id=?',(outcome,'verified' if outcome=='sent' else None,o['id']))
    s.execute('DROP TABLE am_briefing_delivery_attempts')
    sch=Scheduler(s,am.items,am.owner,transport,enabled=lambda:True,briefing=am.briefing_payload)
    expected={'cancelled':'not_sent','unknown':'unknown','sent':'verified'}[outcome]
    assert s.rows('SELECT * FROM am_briefing_delivery_attempts')[0]['state']==expected
    assert s.rows('SELECT * FROM am_briefing_deliveries')[0]['occurrence_id']==o['id']
    assert sch.occurrence(o['id'])['request_key']==o['request_key']
    Scheduler(s,am.items,am.owner,transport,enabled=lambda:True,briefing=am.briefing_payload)
    assert len(s.rows('SELECT * FROM am_briefing_delivery_attempts'))==1


@pytest.mark.parametrize('shape',['table','docx-list','markdown-empty','other-heading'])
def test_F08_structural_positive_and_empty_controls(shape):
    if shape in ('table','docx-list'):
        from docx import Document
        doc=Document()
        if shape=='table':
            table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='交付物清单';table.cell(0,1).text='测试报告A'
        else:
            doc.add_heading('交付物清单',1);doc.add_paragraph('1. 测试报告A')
        buf=io.BytesIO();doc.save(buf);source=read_source('report.docx',buf.getvalue(),source_id='S1',origin='synthetic')
    else:
        text='| 交付物清单 |   |\n| --- | --- |' if shape=='markdown-empty' else '交付物清单：\n版本：v1'
        source=read_source('report.md',text.encode(),source_id='S1',origin='synthetic')
    report=inspect_materials([source],{'checks':[{'id':'1','label':'交付物清单','text':'交付物清单'}]})
    if shape in ('table','docx-list'):assert report['已具备'] and not report['缺少']
    else:assert not report['已具备']

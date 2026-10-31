import io
import pytest
from docx import Document
from fs_agent.doc_logic import read_source
from fs_agent.precheck import inspect_materials

REQ={'checks':[{'id':'1','label':'交付物清单','text':'交付物清单','kind':'presence'}]}

def report(cells,merge=None):
    doc=Document();table=doc.add_table(rows=1,cols=len(cells))
    for i,text in enumerate(cells):table.cell(0,i).text=text
    if merge:table.cell(0,merge[0]).merge(table.cell(0,merge[1]))
    content=io.BytesIO();doc.save(content)
    source=read_source('synthetic.docx',content.getvalue(),source_id='S1',origin='synthetic')
    assert source['read_status']=='complete'
    return inspect_materials([source],REQ)

def test_empty_composite_field_does_not_borrow_responsible_person():
    result=report(('交付物清单','','负责人','AR001合成人员'))
    assert not result['已具备']
    assert result['缺少']

def test_explicit_two_cell_value_is_evidence():
    assert report(('交付物清单','测试报告A'))['已具备']

def test_multiple_field_pairs_bind_only_the_adjacent_value():
    result=report(('项目名称','合成项目','交付物清单','测试报告A','负责人','AR001合成人员'))
    assert result['已具备'][0]['refs']==['S1:3','S1:4']

def test_merged_cells_remain_professional_judgment():
    result=report(('交付物清单','','负责人','AR001合成人员'),(0,1))
    assert not result['已具备'] and not result['缺少']
    assert result['需专业判断']
